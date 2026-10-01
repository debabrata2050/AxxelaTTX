import os
from datetime import datetime
from typing import List, Any, Set
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

from core.models import EXCEL_HEADERS
from core.allocator import parse_trade_date


class InstitutionalExcelExporter:
    """
    Renders structured transfer rows into styled openpyxl Workbooks (S.O.L.I.D SRP).
    Matches institutional clearing desk specifications.
    """

    @classmethod
    def export(
        cls,
        rows: List[List[Any]],
        output_xlsx_path: str,
        separator_row_indices: Set[int] = None
    ) -> str:
        if separator_row_indices is None:
            separator_row_indices = set()
        else:
            separator_row_indices = set(separator_row_indices)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # Ensure headers is first row
        if not rows or rows[0] != EXCEL_HEADERS:
            ws.append(EXCEL_HEADERS)
        for r in rows:
            ws.append(r)

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
        cell_font = Font(name="Calibri", size=11, bold=True)
        cell_align = Alignment(horizontal="center", vertical="center")

        max_row = ws.max_row
        max_col = len(EXCEL_HEADERS)

        for r in range(1, max_row + 1):
            if r in separator_row_indices:
                continue

            # Check if all cells in row are empty/None
            is_empty = all(ws.cell(row=r, column=c).value is None for c in range(1, max_col + 1))
            if is_empty:
                continue

            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = cell_align

            # Data row formatting (skip row 1 header)
            if r > 1:
                # Col 6: Contract Expiry (mmm-yy)
                c_exp_cell = ws.cell(row=r, column=6)
                if isinstance(c_exp_cell.value, str):
                    dt = parse_trade_date(c_exp_cell.value)
                    if dt:
                        c_exp_cell.value = dt
                if isinstance(c_exp_cell.value, datetime):
                    c_exp_cell.number_format = "mmm-yy"

                # Col 14: Date (mm-dd-yy)
                date_cell = ws.cell(row=r, column=14)
                if isinstance(date_cell.value, str):
                    dt = parse_trade_date(date_cell.value)
                    if dt:
                        date_cell.value = dt
                if isinstance(date_cell.value, datetime):
                    date_cell.number_format = "mm-dd-yy"

                # Col 15: Expdate (mm-dd-yy)
                exp_cell = ws.cell(row=r, column=15)
                if isinstance(exp_cell.value, str):
                    dt = parse_trade_date(exp_cell.value)
                    if dt:
                        exp_cell.value = dt
                if isinstance(exp_cell.value, datetime):
                    exp_cell.number_format = "mm-dd-yy"

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                if cell.value is not None:
                    val_str = str(cell.value)
                    if isinstance(cell.value, datetime):
                        val_str = "YYYY-MM-DD"
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

        # Ensure output directory exists
        out_dir = os.path.dirname(os.path.abspath(output_xlsx_path))
        if out_dir and not os.path.exists(out_dir):
            os.makedirs(out_dir, exist_ok=True)

        wb.save(output_xlsx_path)
        return output_xlsx_path
