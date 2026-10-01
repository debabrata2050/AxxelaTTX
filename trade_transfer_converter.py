import argparse
import os
import sys
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def parse_date(val):
    if pd.isna(val) or val is None or str(val).strip() == "":
        return None
    val_str = str(val).strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d-%m-%y"):
        try:
            return datetime.strptime(val_str, fmt)
        except ValueError:
            pass
    try:
        dt = pd.to_datetime(val_str, dayfirst=True)
        return dt.to_pydatetime()
    except Exception:
        return None


def clean_val(val):
    if pd.isna(val) or val is None:
        return None
    val_str = str(val).strip()
    return val_str if val_str != "" else None


def flip_side(side):
    side = str(side).strip().upper() if side else ""
    if side == "B":
        return "S"
    elif side == "S":
        return "B"
    return side


def convert_trade_transfer(
    input_csv_path,
    to_account,
    from_account=None,
    custom_date=None,
    output_xlsx_path=None
):
    if not os.path.exists(input_csv_path):
        raise FileNotFoundError(f"Input file not found: {input_csv_path}")

    # Read CSV, skip completely empty lines
    df = pd.read_csv(input_csv_path, dtype=str)
    # Strip whitespace from column names
    df.columns = [c.strip() for c in df.columns]

    # Filter out empty or invalid rows (where essential columns are missing)
    required_cols = ["contractcode", "transactiontype", "qtybalance"]
    existing_req = [c for c in required_cols if c in df.columns]
    if existing_req:
        df = df.dropna(subset=existing_req, how="all")
        # Also drop rows where contractcode is empty or NaN
        if "contractcode" in df.columns:
            df = df[df["contractcode"].str.strip().fillna("") != ""]

    if df.empty:
        raise ValueError(f"No valid trade rows found in {input_csv_path}")

    # Validate presence of account number in OPD file
    existing_accounts = set()
    if "clientaccountnumber" in df.columns:
        existing_accounts.update(df["clientaccountnumber"].dropna().str.strip().unique())
    if "clientnumber" in df.columns:
        existing_accounts.update(df["clientnumber"].dropna().str.strip().unique())
    existing_accounts = {a for a in existing_accounts if a != ""}

    if not existing_accounts:
        raise ValueError(
            f"Validation Error: No account numbers found in '{os.path.basename(input_csv_path)}'. "
            f"Expected 'clientaccountnumber' or 'clientnumber' columns with values."
        )

    acc_map_upper = {a.upper(): a for a in existing_accounts}

    if from_account:
        source_account_input = from_account.strip()
        if source_account_input.upper() not in acc_map_upper:
            # Account is NOT present in OPD file - validation failure
            similar = [a for a in sorted(existing_accounts) if source_account_input.upper() in a.upper()]
            sample = similar[:10] if similar else sorted(list(existing_accounts))[:10]
            raise ValueError(
                f"Validation Error: Account '{source_account_input}' is NOT present in OPD file '{os.path.basename(input_csv_path)}'.\n"
                f"Total accounts in file: {len(existing_accounts)}. Sample available: {', '.join(sample)}"
            )

        # Match exact casing from file
        source_account = acc_map_upper[source_account_input.upper()]

        mask = pd.Series(False, index=df.index)
        if "clientaccountnumber" in df.columns:
            mask |= (df["clientaccountnumber"].str.strip().str.upper() == source_account.upper())
        if "clientnumber" in df.columns:
            mask |= (df["clientnumber"].str.strip().str.upper() == source_account.upper())
        df = df[mask]
        if df.empty:
            raise ValueError(f"Validation Error: No trades found for account '{source_account}' in {input_csv_path}")
    else:
        if len(existing_accounts) == 1:
            source_account = next(iter(existing_accounts))
            mask = pd.Series(False, index=df.index)
            if "clientaccountnumber" in df.columns:
                mask |= (df["clientaccountnumber"].str.strip() == source_account)
            if "clientnumber" in df.columns:
                mask |= (df["clientnumber"].str.strip() == source_account)
            df = df[mask]
        else:
            sample_accs = sorted(list(existing_accounts))[:8]
            raise ValueError(
                f"Validation Error: File contains multiple accounts ({len(existing_accounts)} accounts found: {', '.join(sample_accs)}...). "
                f"Please specify which account to transfer using --from <account>."
            )

    # Clientgroup for filename
    client_group = "SYM"
    if "clientgroup" in df.columns:
        grp = df["clientgroup"].dropna().str.strip()
        grp = grp[grp != ""]
        if not grp.empty:
            client_group = grp.iloc[0]

    # Determine transfer date
    trade_dt = None
    if custom_date:
        trade_dt = parse_date(custom_date)
    
    if not trade_dt and "datestr" in df.columns:
        d_val = df["datestr"].dropna().str.strip()
        d_val = d_val[d_val != ""]
        if not d_val.empty:
            trade_dt = parse_date(d_val.iloc[0])

    if not trade_dt:
        trade_dt = datetime.now()

    date_compact = trade_dt.strftime("%Y%m%d")

    # Determine output path if not given
    if not output_xlsx_path:
        base_dir = os.path.dirname(os.path.abspath(input_csv_path))
        filename = f"{client_group}.Transfer.{date_compact}({source_account}).xlsx"
        output_xlsx_path = os.path.join(base_dir, filename)

    # Headers for Excel
    headers = [
        "TradeType", "TrdTyp", "TrdSubtyp", "Market ID", "Symbol",
        "Contract Expiry", "TransactionTyp", "Qty", "TradePrice", "Product Type",
        "C/P", "Strike", "Account", "Date", "Expdate",
        "Currency", "Comment", "Fees"
    ]

    # Function to build a row
    def build_row(row_dict, account, side):
        sectyp = clean_val(row_dict.get("sectyp", ""))
        is_opt = (sectyp or "").upper() in ("OPT", "OPTION", "OPTIONS")

        # C/P and Strike only for options
        cp_val = clean_val(row_dict.get("cp", "")) if is_opt else None
        
        strike_val = None
        if is_opt:
            s_raw = clean_val(row_dict.get("strike", ""))
            if s_raw:
                try:
                    s_num = float(s_raw)
                    strike_val = int(s_num) if s_num.is_integer() else s_num
                except ValueError:
                    strike_val = s_raw

        # Expiry date parsing
        contract_expiry = parse_date(row_dict.get("contractexpiry", ""))
        exp_date = parse_date(row_dict.get("expirydate", ""))

        # Market ID (exchangecode)
        market_id = clean_val(row_dict.get("exchangecode", ""))
        if market_id:
            try:
                market_id = int(market_id)
            except ValueError:
                pass

        # Qty
        qty = clean_val(row_dict.get("qtybalance", ""))
        if qty:
            try:
                qty_f = float(qty)
                qty = int(qty_f) if qty_f.is_integer() else qty_f
            except ValueError:
                pass

        # Price
        price = clean_val(row_dict.get("price", ""))
        if price:
            try:
                price = float(price)
            except ValueError:
                pass

        currency = clean_val(row_dict.get("currency", ""))

        return [
            "INTERNAL",                                # TradeType
            clean_val(row_dict.get("trdtyp", "T")),    # TrdTyp
            clean_val(row_dict.get("trdsubtyp", "TRF")),# TrdSubtyp
            market_id,                                 # Market ID
            clean_val(row_dict.get("contractcode", "")),# Symbol
            contract_expiry,                           # Contract Expiry
            side,                                      # TransactionTyp
            qty,                                       # Qty
            price,                                     # TradePrice
            sectyp,                                    # Product Type
            cp_val,                                    # C/P
            strike_val,                                # Strike
            account,                                   # Account
            trade_dt,                                  # Date
            exp_date,                                  # Expdate
            currency,                                  # Currency
            "Transfer",                                # Comment
            "Charge"                                   # Fees
        ]

    # Create Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Styles matching reference Excel
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    cell_font = Font(name="Calibri", size=11, bold=True)
    cell_align = Alignment(horizontal="center", vertical="center")

    # Write header
    ws.append(headers)

    # Leg 1: Destination Account (original side from OPD file)
    for _, row in df.iterrows():
        orig_side = clean_val(row.get("transactiontype", ""))
        row_data = build_row(row, to_account, orig_side)
        ws.append(row_data)

    # Empty separator row (keep track of row index to avoid styling it)
    ws.append([None] * len(headers))
    separator_row_idx = ws.max_row

    # Leg 2: Source Account (reversed side)
    for _, row in df.iterrows():
        orig_side = clean_val(row.get("transactiontype", ""))
        rev_side = flip_side(orig_side)
        row_data = build_row(row, source_account, rev_side)
        ws.append(row_data)

    # Apply font, border, alignment, and number formatting to non-empty rows
    max_row = ws.max_row
    for r in range(1, max_row + 1):
        if r == separator_row_idx:
            continue
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = cell_align

        # Number formatting on data rows
        if r > 1:
            cell_c_exp = ws.cell(row=r, column=6)
            if isinstance(cell_c_exp.value, datetime):
                cell_c_exp.number_format = "mmm-yy"

            cell_date = ws.cell(row=r, column=14)
            if isinstance(cell_date.value, datetime):
                cell_date.number_format = "mm-dd-yy"

            cell_exp = ws.cell(row=r, column=15)
            if isinstance(cell_exp.value, datetime):
                cell_exp.number_format = "mm-dd-yy"

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                if isinstance(cell.value, datetime):
                    val_str = "YYYY-MM-DD"
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    # Save
    wb.save(output_xlsx_path)
    print(f"Generated transfer file: {output_xlsx_path}")
    return output_xlsx_path


def main():
    parser = argparse.ArgumentParser(
        description="Convert OPD Trade CSV to Trade Transfer Excel file."
    )
    parser.add_argument("input_csv", nargs="?", help="Path to input OPD CSV file")
    parser.add_argument("--to", "--to-account", dest="to_account", help="Destination account (e.g. EEABC)")
    parser.add_argument("--from", "--from-account", dest="from_account", help="Source account (override clientaccountnumber)")
    parser.add_argument("--date", dest="custom_date", help="Custom transfer date (YYYY-MM-DD or DD-MM-YYYY)")
    parser.add_argument("--output", "-o", dest="output_xlsx", help="Custom output Excel file path")

    args = parser.parse_args()

    # Interactive prompts if arguments omitted
    input_csv = args.input_csv
    if not input_csv:
        input_csv = input("Enter path to input OPD CSV file: ").strip().strip('"').strip("'")
        if not input_csv:
            print("Error: Input file path required.")
            sys.exit(1)

    to_account = args.to_account
    if not to_account:
        to_account = input("Enter Destination Account (--to, e.g. EEABC): ").strip()
        if not to_account:
            print("Error: Destination account (--to) required.")
            sys.exit(1)

    from_account = args.from_account
    if not from_account and len(sys.argv) <= 1:
        ans = input("Enter Source Account [--from, leave blank to use file clientaccountnumber]: ").strip()
        if ans:
            from_account = ans

    custom_date = args.custom_date
    if not custom_date and len(sys.argv) <= 1:
        ans = input("Enter Transfer Date [leave blank to use file datestr]: ").strip()
        if ans:
            custom_date = ans

    try:
        convert_trade_transfer(
            input_csv_path=input_csv,
            to_account=to_account,
            from_account=from_account,
            custom_date=custom_date,
            output_xlsx_path=args.output_xlsx
        )
    except Exception as e:
        print(f"Error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
