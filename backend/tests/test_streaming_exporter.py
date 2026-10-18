import os
import time
import tempfile
import openpyxl
from datetime import datetime
from core.models import EXCEL_HEADERS
from core.exporter import InstitutionalExcelExporter


def test_streaming_exporter_schema_compliance():
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)

    sample_rows = [
        EXCEL_HEADERS,
        [
            "INTERNAL", "T", "TRF", "NSE", "NIFTY_FUT", datetime(2026, 10, 29),
            "BUY", 100, 24850.50, "FUT", None, None, "ACC_HEDGE_01",
            datetime(2026, 10, 18), datetime(2026, 10, 29), "USD", "TTX Allocation", 0.0
        ]
    ]

    try:
        InstitutionalExcelExporter.export(sample_rows, path)
        assert os.path.exists(path)
        assert os.path.getsize(path) > 0

        # Read back with openpyxl to verify schema
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        assert ws.title == "Sheet1"

        # Check headers match 18-column spec exactly
        read_headers = [ws.cell(row=1, column=c).value for c in range(1, len(EXCEL_HEADERS) + 1)]
        assert read_headers == EXCEL_HEADERS

        # Check row 2 values
        assert ws.cell(row=2, column=1).value == "INTERNAL"
        assert ws.cell(row=2, column=5).value == "NIFTY_FUT"
        assert ws.cell(row=2, column=7).value == "BUY"
        assert ws.cell(row=2, column=8).value == 100
        assert ws.cell(row=2, column=13).value == "ACC_HEDGE_01"
        wb.close()
    finally:
        if os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass


def test_streaming_exporter_scale_benchmark():
    """Benchmark: Streaming 10,000 rows (180,000 cells) must complete in < 10s with low memory."""
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)

    row_template = [
        "INTERNAL", "T", "TRF", "NSE", "NIFTY_FUT", datetime(2026, 10, 29),
        "BUY", 50, 24850.5, "FUT", None, None, "ACC_STREAM_01",
        datetime(2026, 10, 18), datetime(2026, 10, 29), "USD", "Benchmark", 0.0
    ]

    rows = [EXCEL_HEADERS] + [list(row_template) for _ in range(10_000)]

    t0 = time.time()
    try:
        InstitutionalExcelExporter.export(rows, path)
        elapsed = time.time() - t0
        assert os.path.exists(path)
        assert os.path.getsize(path) > 50_000  # Valid XLSX archive
        assert elapsed < 10.0, f"Streaming export too slow: {elapsed:.2f}s (expected < 10.0s)"
    finally:
        if os.path.exists(path):
            try:
                os.remove(path)
            except OSError:
                pass
