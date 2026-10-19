import os
import tempfile
import pytest
import openpyxl

from app import create_app
from core.kgi_netting_engine import KGINettingEngine, DEFAULT_KGI_TO_ACCOUNT, EXCEL_HEADERS

REAL_KGI_CSV = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "data",
    "KGI netting",
    "KT513_20261009_OP (1).csv",
)


@pytest.fixture
def synthetic_csv():
    """Generates synthetic KGI CSV with balanced and unbalanced contracts."""
    content = """TRADE DATE,EXCHANGE,CONTRACT,DEL.MONTH,BUYSELL,QUANTITY,EXPIRYDATE,C.P.F,STRIKE,TRADE PRICE,SETTLEMENT PRICE,UNREALISED PL,PLCCY,DEAL REF,CLIENT CODE
04/06/2026,SGX,TF,202610,B,5,Oct 01 2026 12:00AM,F,0.0000,216.4000,220.0000,100.00,USD,99153778,KT513A23
05/06/2026,SGX,TF,202610,S,5,Oct 01 2026 12:00AM,F,0.0000,218.0000,220.0000,-100.00,USD,99153779,KT513A21
10/06/2026,SGX,XYZ,202611,B,10,Nov 15 2026 12:00AM,F,0.0000,100.0000,100.0000,0.00,USD,99153780,KT513B01
11/06/2026,SGX,XYZ,202611,S,4,Nov 15 2026 12:00AM,F,0.0000,102.0000,100.0000,0.00,USD,99153781,KT513B02
"""
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write(content)
        path = f.name
    yield path
    if os.path.exists(path):
        os.remove(path)


def test_kgi_missing_headers():
    bad_content = "ColA,ColB,ColC\n1,2,3\n"
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write(bad_content)
        path = f.name
    try:
        with pytest.raises(ValueError, match="Missing required columns"):
            KGINettingEngine(path)
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_kgi_empty_file():
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        path = f.name
    try:
        with pytest.raises(ValueError, match="0 bytes"):
            KGINettingEngine(path)
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_kgi_netting_synthetic(synthetic_csv):
    engine = KGINettingEngine(synthetic_csv)
    summary = engine.parse_summary_and_contracts()

    assert summary["total_positions"] == 4
    assert summary["total_products"] == 2
    assert summary["eligible_contracts_count"] == 1
    assert summary["total_eligible_lots"] == 5

    tf_contract = next(c for c in summary["contracts"] if c["contract"] == "TF")
    assert tf_contract["is_eligible"] is True
    assert tf_contract["buy_lots"] == 5
    assert tf_contract["sell_lots"] == 5
    assert tf_contract["net_delta"] == 0

    xyz_contract = next(c for c in summary["contracts"] if c["contract"] == "XYZ")
    assert xyz_contract["is_eligible"] is False
    assert xyz_contract["buy_lots"] == 10
    assert xyz_contract["sell_lots"] == 4
    assert xyz_contract["net_delta"] == 6


def test_kgi_excel_export_layout_a1(synthetic_csv):
    engine = KGINettingEngine(synthetic_csv)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = f.name

    try:
        engine.export_consolidated_excel(
            output_xlsx_path=out_path,
            contracts=["TF"],
            custom_comments={"1": "Batch Netting Test"},
        )

        wb = openpyxl.load_workbook(out_path)
        ws = wb.active
        assert ws.title == "Sheet1"

        # Check Cell A1 layout with zero gaps
        assert ws["A1"].value == "Trade Transfer Form"
        assert ws["A1"].font.name == "Calibri"
        assert ws["A1"].font.bold is True
        assert ws["A1"].fill.start_color.rgb == "00FFFF00"
        assert "<MergedCellRange A1:L1>" in str(ws.merged_cells.ranges)

        # Headers on Row 2
        headers = [ws.cell(2, col).value for col in range(1, 13)]
        assert headers == EXCEL_HEADERS

        # Data on Row 3 (Buy row first)
        assert ws.cell(3, 1).value == 1  # Sl.No.
        assert ws.cell(3, 2).value == "KT513A23"  # FROM
        assert ws.cell(3, 3).value == DEFAULT_KGI_TO_ACCOUNT  # TO: KT513Z
        assert ws.cell(3, 4).value == "TF"  # Instr.
        assert ws.cell(3, 6).value == "B"  # Buy/Sell
        assert ws.cell(3, 7).value == 5  # Lots
        assert ws.cell(3, 12).value == "Batch Netting Test"  # Comments

        # Data on Row 4 (Sell row second)
        assert ws.cell(4, 1).value == 2
        assert ws.cell(4, 2).value == "KT513A21"
        assert ws.cell(4, 6).value == "S"
        assert ws.cell(4, 7).value == 5
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_kgi_real_dataset():
    if not os.path.exists(REAL_KGI_CSV):
        pytest.skip(f"Real data file not present at {REAL_KGI_CSV}")

    engine = KGINettingEngine(REAL_KGI_CSV)
    summary = engine.parse_summary_and_contracts()

    assert summary["total_positions"] == 5617
    assert summary["total_products"] == 34
    assert summary["eligible_contracts_count"] == 9
    assert summary["total_eligible_lots"] == 90

    # Test preview rows for BTP
    preview = engine.get_preview_rows(contracts=["BTP"])
    assert len(preview) == 26
    # Verify B rows come first, then S rows
    buy_rows = [r for r in preview if r["buy_sell"] == "B"]
    sell_rows = [r for r in preview if r["buy_sell"] == "S"]
    assert len(buy_rows) == 13
    assert len(sell_rows) == 13
    assert preview[0]["buy_sell"] == "B"
    assert preview[-1]["buy_sell"] == "S"


def test_kgi_flask_endpoints(synthetic_csv):
    app = create_app()
    client = app.test_client()

    # 1. Upload endpoint via filepath JSON
    resp = client.post("/api/kgi/upload", json={"filepath": synthetic_csv})
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["success"] is True
    assert data["data"]["eligible_contracts_count"] == 1

    # 2. Preview endpoint
    resp = client.post(
        "/api/kgi/preview",
        json={
            "filepath": synthetic_csv,
            "contracts": ["TF"],
            "custom_comments": {"1": "API Preview Comment"},
        },
    )
    assert resp.status_code == 200
    pdata = resp.get_json()
    assert pdata["success"] is True
    assert pdata["count"] == 2
    assert pdata["rows"][0]["comments"] == "API Preview Comment"

    # 3. Export endpoint
    resp = client.post(
        "/api/kgi/export",
        json={
            "filepath": synthetic_csv,
            "contracts": ["TF"],
            "custom_comments": {"1": "API Export Comment"},
            "filename": "Test_KGI_Netting.xlsx",
        },
    )
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.content_type
