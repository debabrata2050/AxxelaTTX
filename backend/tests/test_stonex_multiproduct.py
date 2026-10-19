import os
import pytest
import openpyxl
from core.stonex_netting_engine import StoneXNettingEngine

SAMPLE_CSV_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "data",
    "stoneX netting",
    "NetPositions20261009.csv",
)

@pytest.fixture
def engine():
    if not os.path.exists(SAMPLE_CSV_PATH):
        pytest.skip(f"Sample CSV not found at {SAMPLE_CSV_PATH}")
    return StoneXNettingEngine(SAMPLE_CSV_PATH)

def test_multiproduct_filtering(engine):
    # LGO has 1 net-zero contract (Oct-26) with 546 buy and 546 sell lots (total 1092 rows)
    # NOL has 1 net-zero contract (Nov-26)
    rows = engine.get_preview_rows(symbols=["LGO", "NOL"], eligible_only=True)
    assert len(rows) > 0
    symbols = {r["instr"] for r in rows}
    assert symbols.issubset({"LGO", "NOL"})
    # Verify TO account is constant LCE30553
    assert all(r["to_account"] == "LCE30553" for r in rows)

def test_multicontract_filtering(engine):
    summary = engine.parse_summary_and_contracts()
    eligible_contracts = [c for c in summary["contracts"] if c["is_eligible"]]
    assert len(eligible_contracts) >= 2

    cid1 = eligible_contracts[0]["contract_id"]
    cid2 = eligible_contracts[1]["contract_id"]

    rows = engine.get_preview_rows(contract_ids=[cid1, cid2], eligible_only=True)
    assert len(rows) > 0
    # Rows should only match the two selected contracts
    found_symbols = {r["instr"] for r in rows}
    expected_symbols = {eligible_contracts[0]["symbol"], eligible_contracts[1]["symbol"]}
    assert found_symbols.issubset(expected_symbols)

def test_multiproduct_export(engine, tmp_path):
    out_file = str(tmp_path / "multi_export.xlsx")
    engine.export_consolidated_excel(out_file, symbols=["LGO", "NOL"], eligible_only=True)

    assert os.path.exists(out_file)
    wb = openpyxl.load_workbook(out_file)
    assert "StoneX Netting" in wb.sheetnames
    ws = wb["StoneX Netting"]
    # Verify no yellow fills in cell styles
    for row in ws.iter_rows(min_row=7, max_row=min(20, ws.max_row), min_col=2, max_col=11):
        for cell in row:
            if cell.fill and cell.fill.fill_type:
                assert cell.fill.start_color.rgb not in ["FFFFFF00", "FFFF00"]
