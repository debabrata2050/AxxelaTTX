import os
import tempfile
import pytest
import openpyxl

from app import create_app
from core.stonex_netting_engine import StoneXNettingEngine, DEFAULT_TO_ACCOUNT, EXCEL_HEADERS
from stonex_netting_converter import convert_stonex_netting


@pytest.fixture
def sample_csv():
    """Generates a small synthetic StoneX CSV with both balanced and unbalanced contracts."""
    content = """Grouping,Business Date,Account Id,Family Group,Legal Account Name,Exchange Id,Currency,Exchange Clearing Id,Reporting Symbol,Instrument Description,Trade Type,Option Syle,Strike Price,Quantity,Delivery Period,Net Long Quantity,Net Short Quantity,Weighted Average Price,Market Price,MtM,Market Value,Option Value,Intrinsic Value,Composite Delta,Maturity Date,OTC Indicator,Contract Size,Delivery Cycle,Last Traded Date,First Notice Date,Legal Entity Identifier,FX Rate
ACC01 - SYM1 - Future - Oct-26 - 0.0,2026-10-09,ACC01,GRP1,SYM LTD,CBT,USD,C,SYM1,PRODUCT ONE,Future,,0.0,10,Oct-26,,,,100.0,105.0,,,,,,,,,,,,
ACC02 - SYM1 - Future - Oct-26 - 0.0,2026-10-09,ACC02,GRP1,SYM LTD,CBT,USD,C,SYM1,PRODUCT ONE,Future,,0.0,-10,Oct-26,,,,100.0,105.0,,,,,,,,,,,,
ACC01 - SYM2 - Future - Nov-26 - 0.0,2026-10-09,ACC01,GRP1,SYM LTD,CBT,USD,C,SYM2,PRODUCT TWO,Future,,0.0,5,Nov-26,,,,200.0,210.0,,,,,,,,,,,,
ACC02 - SYM2 - Future - Nov-26 - 0.0,2026-10-09,ACC02,GRP1,SYM LTD,CBT,USD,C,SYM2,PRODUCT TWO,Future,,0.0,-3,Nov-26,,,,200.0,210.0,,,,,,,,,,,,
ACC03 - SYM3 - Future - Dec-26 - 0.0,2026-10-09,ACC03,GRP1,SYM LTD,CBT,USD,C,SYM3,PRODUCT THREE,Future,,0.0,0,Dec-26,,,,50.0,55.0,,,,,,,,,,,,
"""
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write(content)
        path = f.name
    yield path
    if os.path.exists(path):
        os.remove(path)


def test_stonex_missing_headers():
    bad_content = "ColA,ColB,ColC\n1,2,3\n"
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        f.write(bad_content)
        path = f.name
    try:
        with pytest.raises(ValueError, match="Missing required column"):
            StoneXNettingEngine(path)
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_stonex_empty_file():
    with tempfile.NamedTemporaryFile(mode="w", suffix=".csv", delete=False, encoding="utf-8") as f:
        path = f.name
    try:
        with pytest.raises(ValueError, match="0 bytes"):
            StoneXNettingEngine(path)
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_stonex_netting_evaluation(sample_csv):
    engine = StoneXNettingEngine(sample_csv)
    summary = engine.parse_summary_and_contracts()

    assert summary["total_positions"] == 5
    assert summary["total_products"] == 2
    assert summary["eligible_contracts_count"] == 1
    assert summary["total_eligible_lots"] == 10.0

    # SYM1 Oct-26 should be balanced: 10 Buy, 10 Sell -> Net 0
    sym1 = next(c for c in summary["contracts"] if c["symbol"] == "SYM1")
    assert sym1["is_eligible"] is True
    assert sym1["buy_lots"] == 10.0
    assert sym1["sell_lots"] == 10.0
    assert sym1["net_delta"] == 0.0
    assert "NET ZERO" in sym1["status_badge"]

    # SYM2 Nov-26 is unbalanced: 5 Buy, 3 Sell -> Net +2
    sym2 = next(c for c in summary["contracts"] if c["symbol"] == "SYM2")
    assert sym2["is_eligible"] is False
    assert sym2["buy_lots"] == 5.0
    assert sym2["sell_lots"] == 3.0
    assert sym2["net_delta"] == 2.0
    assert "UNBALANCED" in sym2["status_badge"]


def test_stonex_excel_consolidation_single_sheet(sample_csv):
    engine = StoneXNettingEngine(sample_csv)
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = f.name

    try:
        engine.export_consolidated_excel(out_path, eligible_only=True)
        assert os.path.exists(out_path)

        wb = openpyxl.load_workbook(out_path)
        assert len(wb.sheetnames) == 1
        ws = wb["StoneX Netting"]

        # Check Row 1 Title (starts at A1)
        assert ws["A1"].value == "Trade Transfer Form"

        # Check Row 2 Headers
        headers = [ws.cell(row=2, column=c).value for c in range(1, 14)]
        assert headers == EXCEL_HEADERS

        # Check Data Rows (2 rows for SYM1 Oct-26 at rows 3 and 4)
        assert ws.max_row == 4

        row1 = [ws.cell(row=3, column=c).value for c in range(1, 14)]
        assert row1[0] == 1  # Sl.No.
        assert row1[1] == "ACC01"  # FROM
        assert row1[2] == DEFAULT_TO_ACCOUNT  # TO
        assert row1[3] == "SYM1"  # Instr
        assert row1[4] == "PRODUCT ONE"  # Desp
        assert row1[5] == "Buy"  # Side
        assert row1[6] == 10  # Lots (positive)

        row2 = [ws.cell(row=4, column=c).value for c in range(1, 14)]
        assert row2[0] == 2  # Sl.No.
        assert row2[1] == "ACC02"
        assert row2[2] == DEFAULT_TO_ACCOUNT
        assert row2[5] == "Sell"
        assert row2[6] == 10  # abs(-10) -> positive 10

        # Check Zero Color Fill on all cells
        for r in range(1, 5):
            for c in range(1, 14):
                cell = ws.cell(row=r, column=c)
                assert cell.fill.fill_type is None, f"Cell {cell.coordinate} has unexpected fill: {cell.fill.fill_type}"
        wb.close()
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_stonex_converter_cli(sample_csv):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = f.name

    try:
        res = convert_stonex_netting(sample_csv, out_path, eligible_only=True)
        assert os.path.exists(res)
        wb = openpyxl.load_workbook(res)
        assert wb.sheetnames == ["StoneX Netting"]
        wb.close()
    finally:
        if os.path.exists(out_path):
            os.remove(out_path)


def test_stonex_flask_api(sample_csv):
    app = create_app()
    app.config["TESTING"] = True
    client = app.test_client()

    # 1. Test Upload
    with open(sample_csv, "rb") as f:
        resp = client.post("/api/stonex/upload", data={"file": (f, "test.csv")}, content_type="multipart/form-data")
    assert resp.status_code == 200
    res = resp.get_json()
    assert res["success"] is True
    assert res["data"]["eligible_contracts_count"] == 1
    uploaded_filepath = res["data"]["filepath"]

    # 2. Test Preview
    resp2 = client.post("/api/stonex/preview", json={"filepath": uploaded_filepath, "eligible_only": True})
    assert resp2.status_code == 200
    res2 = resp2.get_json()
    assert res2["success"] is True
    assert res2["count"] == 2

    # 3. Test Export with custom_comments
    resp3 = client.post(
        "/api/stonex/export",
        json={
            "filepath": uploaded_filepath,
            "eligible_only": True,
            "custom_comments": {"1": "Transfer note 1", "2": "Netting note 2"},
        },
    )
    assert resp3.status_code == 200
    assert resp3.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        f.write(resp3.data)
        out_excel = f.name
    try:
        wb = openpyxl.load_workbook(out_excel)
        ws = wb["StoneX Netting"]
        assert ws.cell(row=3, column=13).value == "Transfer note 1"
        assert ws.cell(row=4, column=13).value == "Netting note 2"
        wb.close()
    finally:
        if os.path.exists(out_excel):
            os.remove(out_excel)
