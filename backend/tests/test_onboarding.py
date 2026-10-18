import os
import pytest
import openpyxl

from core.onboarding_engine import (
    derive_clientgroup,
    derive_clientsubgroup,
    derive_broker,
    build_row,
    derive_preview,
    build_onboarding_excel,
    _sanitize_cell,
    COLUMN_ORDER,
)
from core.onboarding_db import (
    list_clientgroup_rules,
    upsert_clientgroup_rule,
    delete_clientgroup_rule,
    list_broker_rules,
    upsert_broker_rule,
    delete_broker_rule,
    list_subgroup_suffixes,
    upsert_subgroup_suffix,
    delete_subgroup_suffix,
    save_sheet_records_cache,
    search_cached_accounts,
    lookup_cached_account,
    batch_lookup_cached_accounts,
    get_cached_records_count,
)
from core.sheet_service import (
    clean_title_case,
    split_joint_str,
    parse_names,
    normalize_location,
    normalize_sub_branch,
    extract_sheet_id_and_gid,
    get_sheet_service,
)


# ==============================================================================
# Onboarding Engine Tests
# ==============================================================================

def test_onboarding_engine_clientgroup_derivation():
    # ALGO comms code has highest priority
    assert derive_clientgroup("ALGO", "Kolkata", "Senior", False) == "ALGO"

    # Kolkata location with sub_branch
    assert derive_clientgroup("CWSYM", "Kolkata", "Senior", False) == "KOLSR"
    assert derive_clientgroup("CWSYM", "Kolkata", "Junior", False) == "KOLJR"

    # Gurgaon with commodity vs non-commodity
    assert derive_clientgroup("CWSYM", "Gurgaon", None, True) == "GURCOM"
    assert derive_clientgroup("CWSYM", "Gurgaon", None, False) == "GURFI"

    # Regional branches
    assert derive_clientgroup("CWSYM", "Bengaluru", None, False) == "BAN"
    assert derive_clientgroup("CWSYM", "Mumbai", None, False) == "MUM"
    assert derive_clientgroup("CWSYM", "Dubai", None, False) == "GLB"


def test_onboarding_engine_clientsubgroup_derivation():
    # Known mapping with 8-char prefix
    assert derive_clientsubgroup("KOLSR", "20261018") == "20261018KOLSR"
    assert derive_clientsubgroup("BAN", "20261018") == "20261018BLORE"

    # Prefix longer than 8 chars is clipped
    assert derive_clientsubgroup("GURFI", "202610189999") == "20261018GURFI"

    # No prefix returns pure suffix
    assert derive_clientsubgroup("ALGO", "") == "ALGOX"


def test_onboarding_engine_broker_derivation():
    # Longest prefix matches first: EEGG before EE
    dist, clr = derive_broker("EEGG12345")
    assert dist == "MARIX"
    assert clr == "CLR MARIX SYM"

    dist, clr = derive_broker("EE00123")
    assert dist == "Marex"
    assert clr == "CLR MRX SYM"

    dist, clr = derive_broker("KT9999")
    assert dist == "KGI"
    assert clr == "CLR KGI SYM"

    dist, clr = derive_broker("UNKNOWN_PREFIX")
    assert dist == ""
    assert clr == ""


def test_onboarding_engine_build_row_schema():
    user_input = {
        "firstname": "John",
        "lastname": "Doe",
        "clientid": "EEGG001",
        "email": "john.doe@example.com",
        "commsgroupcode": "CWSYM",
        "location": "Kolkata",
        "sub_branch": "Senior",
        "is_commodity": False,
        "subgroupPrefix": "20261018",
    }
    row = build_row(user_input)
    assert len(row) >= len(COLUMN_ORDER)
    assert row["firstname"] == "John"
    assert row["lastname"] == "Doe"
    assert row["clientid"] == "EEGG001"
    assert row["clientgroup"] == "KOLSR"
    assert row["clientsubgroup"] == "20261018KOLSR"
    assert row["distributor"] == "MARIX"
    assert row["cleareraccountid"] == "CLR MARIX SYM"
    assert row["clientcountrycode"] == "IN"
    assert row["basecurrency"] == "USD"


def test_onboarding_excel_export_and_formula_injection():
    # Dangerous formula injection in input
    malicious_input = [
        {
            "firstname": "=CMD|' /C calc'!A0",
            "lastname": "+12345",
            "clientid": "EEGG999",
            "email": "@admin.leak",
            "commsgroupcode": "ALGO",
            "location": "Mumbai",
        }
    ]

    xlsx_path = build_onboarding_excel(malicious_input)
    try:
        assert os.path.exists(xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path, data_only=False)
        ws = wb.active
        # Header rows 1-3
        assert ws.cell(row=1, column=1).value == "Required YN"
        assert ws.cell(row=2, column=1).value == "Collection"
        assert ws.cell(row=3, column=1).value == "Field Name"

        # Data row starts at row 4
        # First name (col 2, after blank col A) should be sanitized with leading quote
        fn_val = ws.cell(row=4, column=2).value
        assert fn_val.startswith("'=")

        # Email cell (col 9, email column) should be sanitized with leading quote
        email_val = ws.cell(row=4, column=9).value
        assert email_val.startswith("'@")
    finally:
        if os.path.exists(xlsx_path):
            os.remove(xlsx_path)


# ==============================================================================
# Onboarding SQLite Cache & Search Tests
# ==============================================================================

def test_onboarding_db_cache_and_chunked_batch_lookup():
    # Populate test records in cache
    test_records = [
        {
            "account": f"TEST_ACC_{i:04d}",
            "raw_name": f"Trader {i}",
            "location": "Kolkata" if i % 2 == 0 else "Mumbai",
            "sub_branch": "Senior" if i % 2 == 0 else "",
            "subgroup_prefix": "20261018",
            "commsgroupcode": "CWSYM",
            "is_joint": False,
            "persons": [{"label": f"Trader {i}", "input": {}, "derived": {}}],
        }
        for i in range(1200)  # > 1000 items to test chunking
    ]

    save_sheet_records_cache(test_records)
    assert get_cached_records_count() >= 1200

    # Test fast autocomplete search
    results = search_cached_accounts("TEST_ACC_0005", limit=5)
    assert len(results) == 1
    assert results[0]["account"] == "TEST_ACC_0005"

    # Test single lookup
    single = lookup_cached_account("TEST_ACC_0010")
    assert single is not None
    assert single["account"] == "TEST_ACC_0010"

    # Test chunked batch lookup with 1200 keys (avoids SQLite variable limit)
    keys_to_fetch = [f"TEST_ACC_{i:04d}" for i in range(1200)]
    batch_res = batch_lookup_cached_accounts(keys_to_fetch)
    assert len(batch_res) == 1200
    assert "TEST_ACC_1199" in batch_res


# ==============================================================================
# Sheet Service Helpers & Joint Account Splitting Tests
# ==============================================================================

def test_sheet_service_helpers():
    assert clean_title_case("jOhN doe") == "John Doe"
    assert split_joint_str("Alice / Bob") == ["Alice", "Bob"]
    assert split_joint_str("Charlie ; David") == ["Charlie", "David"]

    # Single name parsing
    single = parse_names("Rahul Sharma")
    assert single == [("Rahul", "Sharma")]

    # Joint name parsing
    joint = parse_names("Priya Patel / Amit Patel")
    assert len(joint) == 2
    assert joint[0] == ("Priya", "Patel")
    assert joint[1] == ("Amit", "Patel")

    # Location normalization
    assert normalize_location("gurugram") == "Gurgaon"
    assert normalize_location("bangalore") == "Bengaluru"
    assert normalize_location("kolkata") == "Kolkata"

    # URL parser
    sheet_id, gid = extract_sheet_id_and_gid("https://docs.google.com/spreadsheets/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#gid=123456")
    assert sheet_id == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
    assert gid == 123456
