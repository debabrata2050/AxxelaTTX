import os
import openpyxl
import pytest

from trade_transfer_converter import convert_trade_transfer


def test_converter_cli_single_account_auto_detection(make_trade_csv, tmp_path):
    csv_file = make_trade_csv(num_rows=20, num_accounts=1, file_prefix="cli_single")
    output_xlsx = str(tmp_path / "transfer_out.xlsx")

    # When from_account is not specified and only 1 account exists, it auto-detects
    result_path = convert_trade_transfer(
        input_csv_path=csv_file,
        to_account="DEST_ACC",
        output_xlsx_path=output_xlsx,
    )

    assert os.path.exists(result_path)
    wb = openpyxl.load_workbook(result_path, data_only=True)
    ws = wb.active

    # Check headers
    assert ws.cell(row=1, column=1).value == "TradeType"
    assert ws.cell(row=1, column=5).value == "Symbol"
    assert ws.cell(row=1, column=13).value == "Account"

    # Must contain both Legs (Dest and Source)
    dest_count = 0
    src_count = 0
    for r in range(2, ws.max_row + 1):
        acc = ws.cell(row=r, column=13).value
        if acc == "DEST_ACC":
            dest_count += 1
        elif acc == "ACC_001":
            src_count += 1

    assert dest_count == 20
    assert src_count == 20


def test_converter_cli_multi_account_explicit_routing(make_trade_csv, tmp_path):
    csv_file = make_trade_csv(num_rows=60, num_accounts=3, file_prefix="cli_multi")
    output_xlsx = str(tmp_path / "multi_out.xlsx")

    # Explicit multi-account transfer: ACC_001, ACC_002 -> DEST_001, DEST_002
    result_path = convert_trade_transfer(
        input_csv_path=csv_file,
        from_account="ACC_001,ACC_002",
        to_account="DEST_001,DEST_002",
        output_xlsx_path=output_xlsx,
        price_mode="settle",
    )

    assert os.path.exists(result_path)
    wb = openpyxl.load_workbook(result_path, data_only=True)
    ws = wb.active

    accounts_in_file = {
        ws.cell(row=r, column=13).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=13).value
    }

    assert "DEST_001" in accounts_in_file
    assert "DEST_002" in accounts_in_file
    assert "ACC_001" in accounts_in_file
    assert "ACC_002" in accounts_in_file
    # ACC_003 was not routed, so must not appear
    assert "ACC_003" not in accounts_in_file


def test_converter_cli_contract_filter(make_trade_csv, tmp_path):
    csv_file = make_trade_csv(num_rows=30, num_accounts=1, file_prefix="cli_filter")
    output_xlsx = str(tmp_path / "filter_out.xlsx")

    # Filter to only FUT contracts
    result_path = convert_trade_transfer(
        input_csv_path=csv_file,
        to_account="DEST_ACC",
        filter_product="FUT",
        output_xlsx_path=output_xlsx,
    )

    assert os.path.exists(result_path)
    wb = openpyxl.load_workbook(result_path, data_only=True)
    ws = wb.active

    for r in range(2, ws.max_row + 1):
        prod = ws.cell(row=r, column=10).value  # Product Type
        if prod:
            assert prod == "FUT"


def test_converter_cli_account_not_found_raises(make_trade_csv):
    csv_file = make_trade_csv(num_rows=10, num_accounts=1, file_prefix="cli_err")

    with pytest.raises(ValueError, match="Source account 'NON_EXISTENT' not found"):
        convert_trade_transfer(
            input_csv_path=csv_file,
            from_account="NON_EXISTENT",
            to_account="DEST",
        )
